"""Stamp a project's Project Setup tab in a fresh copy of the master register.

One-off setup helper — run once per project to mint the master register, then
upload it to /Projects/<Suburb> - <Address>/Variations/ in Dropbox.

Example:
    python3 scripts/setup_project_register.py \\
        --project-code "137a" \\
        --project-name "137 Annandale Street, Annandale" \\
        --client "Jon and Stacey Day" \\
        --client-email "jonathon.day@wentworthcp.com.au; stacey@stelladay.com.au" \\
        --contract-date 2025-08-20 \\
        --contract-sum 2997386 \\
        --dropbox-folder "Annandale - 137 Annandale Street" \\
        --out "working/137 Annandale Variation Register.xlsx"
"""
from __future__ import annotations

import argparse
import datetime as dt
import shutil
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = REPO_ROOT / "templates" / "master_register_template.xlsx"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project-code", required=True,
                    help='Short project code used everywhere, e.g. "137a"')
    ap.add_argument("--project-name", required=True,
                    help='Full project name / site address, e.g. "137 Annandale Street, Annandale"')
    ap.add_argument("--client", required=True,
                    help='Client name(s), e.g. "Jon and Stacey Day"')
    ap.add_argument("--client-email", default="",
                    help='Client email(s); multiples separated by "; "')
    ap.add_argument("--client-phone", default="",
                    help="Client phone (optional)")
    ap.add_argument("--contract-date", required=True,
                    help="ISO date YYYY-MM-DD")
    ap.add_argument("--contract-sum", required=True, type=float,
                    help="Original contract sum, inc GST, as a number")
    ap.add_argument("--dropbox-folder", required=True,
                    help='Dropbox project folder name under /Projects, '
                         'e.g. "Annandale - 137 Annandale Street"')
    ap.add_argument("--builder-abn", default="72 098 695 446",
                    help="Builder ABN (default: Gold and Eagle Constructions)")
    ap.add_argument("--builder-licence", default="139701C")
    ap.add_argument("--out", required=True, help="Output xlsx path")
    args = ap.parse_args()

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(TEMPLATE, out)

    wb = load_workbook(out)
    ws = wb["Project Setup"]

    # Map values to the EXACT labels already present in the master template's
    # column A. We write into the existing rows — no new rows appended, no
    # duplicates introduced.
    setup = {
        "Project code": args.project_code,
        "Project name / site address": args.project_name,
        "Client name(s)": args.client,
        "Client email": args.client_email,
        "Client phone": args.client_phone,
        "Contract date": dt.date.fromisoformat(args.contract_date),
        "Original contract sum (inc GST)": args.contract_sum,
        "Builder ABN": args.builder_abn,
        "Builder NSW Contractor Licence #": args.builder_licence,
        "Dropbox root for project": f"/Projects/{args.dropbox_folder}",
    }

    # Build a case-insensitive label -> row map from column A.
    label_rows = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.strip():
            label_rows[v.strip().lower()] = r

    unmatched = []
    for key, val in setup.items():
        row = label_rows.get(key.lower())
        if row is None:
            unmatched.append(key)
            continue
        ws.cell(row=row, column=2, value=val)

    if unmatched:
        print(f"WARNING: these labels were not found in the template "
              f"(skipped, not written): {unmatched}")

    wb.save(out)
    print(f"Created: {out}")
    print(f"  Project code:  {args.project_code}")
    print(f"  Project:       {args.project_name}")
    print(f"  Client:        {args.client}")
    print(f"  Client email:  {args.client_email or '(none)'}")
    print(f"  Contract date: {args.contract_date}")
    print(f"  Contract sum:  ${args.contract_sum:,.2f} inc GST")
    print(f"  Builder ABN:   {args.builder_abn}")
    print(f"  Dropbox root:  /Projects/{args.dropbox_folder}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
